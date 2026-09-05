import json
from pathlib import Path

from docx import Document

from applypilot.scoring.pdf import (
    _record_export_artifacts,
    build_html,
    convert_to_docx,
    parse_resume,
)
from applypilot.scoring.tailor import (
    assemble_resume_text,
    build_diff_report,
    build_unified_diff,
)
from applypilot.scoring.validator import validate_tailored_resume

MASTER_RESUME = """Juno Wang
AI Researcher
juno@example.test | github.com/juno

SUMMARY
AI researcher working on agentic AI and LLM evaluation.

SKILLS
Programming: Python, SQL
AI & LLM: Agentic AI, RAG, LLM Evaluation
Tools: Git, AWS

EXPERIENCE
AI Research Intern at Example Lab
Python | 2025 - Present
- Built an agentic AI evaluation pipeline that improved accuracy by 10%.

PROJECTS
Research Copilot
Python, RAG | 2026
- Built a research retrieval workflow.

PUBLICATIONS
First Author - Agent Evaluation, VLDB, 2023

EDUCATION
University at Buffalo | Ph.D. Candidate, Computer Science | 2023 - Present
"""


PROFILE = {
    "personal": {
        "full_name": "Juno Wang",
        "email": "juno@example.test",
        "github_url": "github.com/juno",
    },
    "resume_facts": {
        "preserved_companies": ["Example Lab"],
        "preserved_projects": ["Research Copilot"],
        "preserved_school": "University at Buffalo",
        "real_metrics": ["10%"],
    },
    "skills_boundary": {
        "programming": ["Python", "SQL"],
        "ai": ["Agentic AI", "RAG", "LLM Evaluation"],
        "tools": ["Git", "AWS"],
    },
}


TAILORED_DATA = {
    "title": "Research Scientist, LLM Agents",
    "summary": "AI researcher focused on agentic AI, RAG, and LLM evaluation.",
    "skills": {
        "AI & LLM": "Agentic AI, LLM Evaluation, RAG",
        "Programming": "Python, SQL",
        "Tools": "AWS, Git",
    },
    "experience": [
        {
            "header": "AI Research Intern at Example Lab",
            "subtitle": "Python | 2025 - Present",
            "bullets": [
                "Built an agentic AI evaluation pipeline and improved accuracy by 10%."
            ],
        }
    ],
    "projects": [
        {
            "header": "Research Copilot",
            "subtitle": "Python, RAG | 2026",
            "bullets": ["Built a research retrieval workflow."],
        }
    ],
    "education": "This model-provided value must be ignored",
}


def _tailored_text() -> str:
    return assemble_resume_text(
        TAILORED_DATA,
        PROFILE,
        original_education=(
            "University at Buffalo | Ph.D. Candidate, Computer Science | 2023 - Present"
        ),
        original_publications="First Author - Agent Evaluation, VLDB, 2023",
    )


def test_source_grounded_validation_accepts_supported_facts() -> None:
    result = validate_tailored_resume(
        _tailored_text(),
        PROFILE,
        original_text=MASTER_RESUME,
        mode="normal",
    )
    assert result["passed"] is True
    assert result["errors"] == []


def test_source_grounded_validation_rejects_new_skill_and_number() -> None:
    altered = _tailored_text().replace("AWS, Git", "AWS, Kubernetes, Git")
    altered = altered.replace("10%", "95%")

    result = validate_tailored_resume(
        altered,
        PROFILE,
        original_text=MASTER_RESUME,
        mode="normal",
    )

    assert result["passed"] is False
    assert any("Kubernetes" in error for error in result["errors"])
    assert any("95%" in error for error in result["errors"])


def test_diff_report_is_auditable() -> None:
    tailored = _tailored_text()
    report = build_diff_report(MASTER_RESUME, tailored)
    diff = build_unified_diff(MASTER_RESUME, tailored)

    assert report["original_sha256"] != report["tailored_sha256"]
    assert report["added_lines"] > 0
    assert "--- master-resume.txt" in diff
    assert "+++ tailored-resume.txt" in diff


def test_docx_is_single_column_and_preserves_text(tmp_path: Path) -> None:
    text_path = tmp_path / "tailored.txt"
    text_path.write_text(_tailored_text(), encoding="utf-8")

    docx_path = convert_to_docx(text_path)
    document = Document(docx_path)
    extracted = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert len(document.sections) == 1
    assert not document.tables
    assert "Research Scientist, LLM Agents" in extracted
    assert "University at Buffalo" in extracted
    assert "First Author - Agent Evaluation" in extracted
    assert docx_path.stat().st_mode & 0o777 == 0o600


def test_html_escapes_resume_content() -> None:
    parsed = parse_resume(_tailored_text().replace("AI researcher", "AI <researcher> & builder"))
    rendered = build_html(parsed)

    assert "AI &lt;researcher&gt; &amp; builder" in rendered
    assert "First Author - Agent Evaluation" in rendered
    assert '<div class="education">University at Buffalo' in rendered
    assert "AI <researcher> & builder" not in rendered


def test_export_paths_are_added_to_audit_report(tmp_path: Path) -> None:
    report_path = tmp_path / "report.json"
    docx_path = tmp_path / "resume.docx"
    pdf_path = tmp_path / "resume.pdf"
    report_path.write_text('{"status": "approved", "artifacts": {}}', encoding="utf-8")

    _record_export_artifacts(str(report_path), docx_path, pdf_path)

    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["artifacts"]["docx"] == str(docx_path)
    assert report["artifacts"]["pdf"] == str(pdf_path)
    assert report_path.stat().st_mode & 0o777 == 0o600
